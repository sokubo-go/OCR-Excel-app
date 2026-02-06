import calendar
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import TimecardData

# スタイル定義
TITLE_FONT = Font(name="Yu Gothic", bold=True, size=14)
SUBTITLE_FONT = Font(name="Yu Gothic", size=9)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(name="Yu Gothic", bold=True, size=9, color="FFFFFF")
CELL_FONT = Font(name="Yu Gothic", size=9)
BOLD_FONT = Font(name="Yu Gothic", bold=True, size=9)
TOTAL_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
SAT_FILL = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
SUN_FILL = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")
HOLIDAY_FILL = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
GREEN_FONT = Font(name="Yu Gothic", bold=True, size=9, color="008000")


def generate_excel(timecard_list: list[TimecardData]) -> io.BytesIO:
    """タイムカードデータをExcelファイルに変換する"""
    wb = Workbook()
    wb.remove(wb.active)

    for idx, timecard in enumerate(timecard_list):
        sheet_name = timecard.employee_name or f"タイムカード{idx + 1}"
        for ch in ["\\", "/", "*", "?", ":", "[", "]"]:
            sheet_name = sheet_name.replace(ch, "")
        sheet_name = sheet_name[:31]
        ws = wb.create_sheet(title=sheet_name)
        _write_timecard_sheet(ws, timecard)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def merge_excels(excel_buffers: list[io.BytesIO]) -> io.BytesIO:
    """複数のExcelファイルを1つにマージする"""
    from openpyxl import load_workbook
    merged_wb = Workbook()
    merged_wb.remove(merged_wb.active)

    for buf in excel_buffers:
        buf.seek(0)
        src_wb = load_workbook(buf)
        for src_ws in src_wb.worksheets:
            # シート名の重複を避ける
            name = src_ws.title
            counter = 1
            while name in merged_wb.sheetnames:
                name = f"{src_ws.title}_{counter}"
                counter += 1
            dest_ws = merged_wb.create_sheet(title=name)
            _copy_sheet(src_ws, dest_ws)

    output = io.BytesIO()
    merged_wb.save(output)
    output.seek(0)
    return output


def _copy_sheet(src_ws, dest_ws):
    """シートの内容とスタイルをコピーする"""
    # 列幅コピー
    for col_letter, dim in src_ws.column_dimensions.items():
        dest_ws.column_dimensions[col_letter].width = dim.width

    # 行高さコピー
    for row_num, dim in src_ws.row_dimensions.items():
        dest_ws.row_dimensions[row_num].height = dim.height

    # マージセルコピー
    for merged_range in src_ws.merged_cells.ranges:
        dest_ws.merge_cells(str(merged_range))

    # セルデータ・スタイルコピー
    for row in src_ws.rows:
        for cell in row:
            dest_cell = dest_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                dest_cell.font = cell.font.copy()
                dest_cell.fill = cell.fill.copy()
                dest_cell.border = cell.border.copy()
                dest_cell.alignment = cell.alignment.copy()
                dest_cell.number_format = cell.number_format


def _write_timecard_sheet(ws, timecard: TimecardData) -> None:
    """スクリーンショットのフォーマットに合わせてシートを書き込む"""

    # 列幅
    widths = {"A": 5, "B": 5, "C": 10, "D": 10, "E": 8, "F": 10, "G": 10, "H": 12, "I": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # --- Row 1: タイトル ---
    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value = f"勤務表（タイムカード）　{timecard.year}年{timecard.month}月分"
    c.font = TITLE_FONT
    c.alignment = CENTER
    ws.row_dimensions[1].height = 30

    # --- Row 2: 契約期間 ---
    ws.merge_cells("A2:I2")
    c = ws["A2"]
    if timecard.contract_period:
        c.value = f"契約期間: {timecard.contract_period}"
    c.font = SUBTITLE_FONT
    c.alignment = CENTER
    ws.row_dimensions[2].height = 18

    # --- Row 3: ヘッダー ---
    headers = ["日", "曜日", "出勤時刻", "退勤時刻", "休憩(分)", "実働時間", "契約外時間", "備考", "管理者確認"]
    header_row = 3
    ws.row_dimensions[header_row].height = 22

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = CENTER

    # --- データ行: 月の全日分 ---
    # recordsを日付でマッピング
    record_map = {}
    for r in timecard.records:
        record_map[r.date.strip()] = r

    year = int(timecard.year) if timecard.year.isdigit() else 2025
    month = int(timecard.month) if timecard.month.isdigit() else 1
    days_in_month = calendar.monthrange(year, month)[1]

    weekday_names = ["月", "火", "水", "木", "金", "土", "日"]

    data_start_row = header_row + 1
    for day in range(1, days_in_month + 1):
        row = data_start_row + day - 1
        ws.row_dimensions[row].height = 18

        weekday_idx = calendar.weekday(year, month, day)
        dow = weekday_names[weekday_idx]

        rec = record_map.get(str(day))

        values = [
            str(day),
            dow,
            rec.clock_in if rec else "",
            rec.clock_out if rec else "",
            rec.break_minutes if rec else "",
            rec.working_hours if rec else "",
            rec.overtime if rec else "",
            rec.remarks if rec else "",
            rec.manager if rec else "",
        ]

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER

        # 土日の色分け
        if dow == "土":
            for col_idx in range(1, 10):
                ws.cell(row=row, column=col_idx).fill = SAT_FILL
        elif dow == "日":
            for col_idx in range(1, 10):
                ws.cell(row=row, column=col_idx).fill = SUN_FILL

        # 有給休暇等の色分け
        if rec and rec.remarks and "有給" in rec.remarks:
            for col_idx in range(1, 10):
                ws.cell(row=row, column=col_idx).fill = HOLIDAY_FILL

    # --- 合計行 ---
    total_row = data_start_row + days_in_month
    ws.row_dimensions[total_row].height = 22

    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    c = ws.cell(row=total_row, column=1, value="合計")
    c.font = BOLD_FONT
    c.fill = TOTAL_FILL
    c.border = THIN_BORDER
    c.alignment = CENTER
    ws.cell(row=total_row, column=2).border = THIN_BORDER
    ws.cell(row=total_row, column=2).fill = TOTAL_FILL

    # 出勤日数（出勤時刻が入っている日数をカウント）
    col_c = get_column_letter(3)  # 出勤時刻列
    data_end_row = total_row - 1
    label_cell = ws.cell(row=total_row, column=3, value="出勤日数:")
    label_cell.font = BOLD_FONT
    label_cell.fill = TOTAL_FILL
    label_cell.border = THIN_BORDER
    label_cell.alignment = CENTER

    count_cell = ws.cell(row=total_row, column=4)
    count_cell.value = f'=COUNTA({col_c}{data_start_row}:{col_c}{data_end_row})'
    count_cell.font = BOLD_FONT
    count_cell.fill = TOTAL_FILL
    count_cell.border = THIN_BORDER
    count_cell.alignment = CENTER

    # 実働合計ラベル
    label2 = ws.cell(row=total_row, column=5, value="実働合計:")
    label2.font = BOLD_FONT
    label2.fill = TOTAL_FILL
    label2.border = THIN_BORDER
    label2.alignment = CENTER

    # 実働合計値（実働時間列のSUM）
    col_f = get_column_letter(6)
    sum_cell = ws.cell(row=total_row, column=6)
    sum_cell.value = f'=SUM({col_f}{data_start_row}:{col_f}{data_end_row})'
    sum_cell.font = GREEN_FONT
    sum_cell.fill = TOTAL_FILL
    sum_cell.border = THIN_BORDER
    sum_cell.alignment = CENTER
    sum_cell.number_format = '0.00'

    for col_idx in [7, 8, 9]:
        cell = ws.cell(row=total_row, column=col_idx)
        cell.fill = TOTAL_FILL
        cell.border = THIN_BORDER

    # --- 集計サマリー ---
    summary_start = total_row + 2
    ws.merge_cells(start_row=summary_start, start_column=1, end_row=summary_start, end_column=4)
    c = ws.cell(row=summary_start, column=1, value="【集計サマリー】")
    c.font = BOLD_FONT

    summary_items = [
        ("出勤日数", f"=COUNTA({col_c}{data_start_row}:{col_c}{data_end_row})", "日"),
        ("有給休暇", f'=COUNTIF(H{data_start_row}:H{data_end_row},"*有給*")', "日"),
        ("合計実働時間", f"=SUM({col_f}{data_start_row}:{col_f}{data_end_row})", "時間"),
        (
            "1日平均実働",
            f'=IF(COUNTA({col_c}{data_start_row}:{col_c}{data_end_row})>0,'
            f'SUM({col_f}{data_start_row}:{col_f}{data_end_row})/COUNTA({col_c}{data_start_row}:{col_c}{data_end_row}),0)',
            "時間",
        ),
    ]

    for i, (label, formula, unit) in enumerate(summary_items):
        row = summary_start + 1 + i
        ws.cell(row=row, column=1, value=label).font = CELL_FONT
        val_cell = ws.cell(row=row, column=2, value=formula)
        val_cell.font = BOLD_FONT
        val_cell.alignment = CENTER
        val_cell.number_format = '0.00' if "時間" in unit else '0'
        ws.cell(row=row, column=3, value=unit).font = CELL_FONT
